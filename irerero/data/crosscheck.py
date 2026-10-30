"""Received files that verify the staged tables instead of replacing them.

Each function reads one kind of received file, compares it with what is already staged in
data/raw/<dataset>/, and returns intake checks. Nothing here writes to data/raw/: a file that
is national only, in five-year age groups or at district level cannot stand in for the sector
and single-age tables the app reads, and no value is spread down from it.

    PHC5-2022_Main_Indicators.xlsx           census sector populations, district households,
                                             water, sanitation and electricity
    2.7 Muhanga_single_age.xlsx              the Muhanga rows of the staged projections
    <Province>_5_yrs_age.xlsx                province totals of the staged projections
    Rwanda_MINEDUC_PrePrimary_*.xlsx         national totals of the staged district yearbook
    rwanda_preprimary_schools.csv            the same records as the school GeoJSON
    WorldPop *.tif                           which age and sex layers arrived, and their grid
"""

import re
from pathlib import Path

import numpy as np
import pandas as pd

from irerero.config import RAW, REFERENCE
from irerero.data.crosswalk import district_ids, sector_ids
from irerero.data.sources import census, projections, worldpop, yearbook
from irerero.data.validate import SchemaError

# Checks share the intake's Check record; imported lazily to avoid a cycle.
PHC5_SHEETS = {"Contents", "Table 1", "Table 96", "Table 100"}
PHC5_SECTOR_TABLES = ("Table 96", "Table 97", "Table 98", "Table 99", "Table 100")
PHC5 = "PHC5 2022 Main Indicators"
MINEDUC_SHEETS = {"Annual_Summary", "Long_Data"}
PROVINCE_CODES = {
    "city of kigali": "RW1",
    "southern province": "RW2",
    "western province": "RW3",
    "northern province": "RW4",
    "eastern province": "RW5",
}
SCHOOL_CSV_COLUMNS = {"school_code", "school_category", "geometry_latitude", "geometry_longitude"}
# WorldPop layers: sex, age group (0 = under 1, 1 = 1 to 4) and year, in both naming schemes.
WORLDPOP_NAME = re.compile(
    r"^rwa_(?P<sex>[fmT])_(?P<age>\w+?)_(?P<year>\d{4})_(?P<product>constrained_UNadj|CN_100m_R2025A_v1)\.tif$"
)
UNDER5 = {"0", "1", "00", "01"}


def _check(dataset: str, level: str, text: str):  # type: ignore[no-untyped-def]
    from irerero.data.intake import Check

    return Check(dataset, level, text)


def _text(value: object) -> str:
    return " ".join(str(value).replace("\xa0", " ").split())


def _staged(raw: Path, key: str, file: str) -> pd.DataFrame | None:
    path = raw / key / file
    return pd.read_csv(path) if path.exists() else None


# --- PHC5 Main Indicators -----------------------------------------------------------------
def is_phc5(sheets: list[str]) -> bool:
    return set(sheets) >= PHC5_SHEETS


def _district_table(sheet) -> pd.Series:  # type: ignore[no-untyped-def]
    """First numeric column of a district table, by printed name (Rwanda and provinces too)."""
    out = {}
    for row in sheet.iter_rows(values_only=True):
        if isinstance(row[1], str) and isinstance(row[2], int | float):
            out[_text(row[1])] = row[2]
    return pd.Series(out, dtype=float)


def phc5_sector_population(workbook) -> pd.DataFrame:  # type: ignore[no-untyped-def]
    """Tables 96 to 100: one row per sector (district, sector, total population)."""
    rows = []
    for name in PHC5_SECTOR_TABLES:
        district = None
        for row in list(workbook[name].iter_rows(values_only=True))[4:]:
            label, total = row[1], row[2]
            if not isinstance(label, str) or not isinstance(total, int | float):
                continue
            if label.startswith(" "):  # sectors are indented under their district
                rows.append((district, _text(label), int(total)))
            else:
                district = _text(label)
    return pd.DataFrame(rows, columns=["district_name", "sector_name", "population"])


def phc5_main_indicators(path: Path, raw: Path = RAW) -> list:
    """Census sector populations and district household tables against the staged census."""
    from openpyxl import load_workbook

    key = census.DATASET.key
    workbook = load_workbook(path, read_only=True, data_only=True)
    checks = [_check(key, "pass", f"{PHC5}: {len(workbook.sheetnames)} sheets opened")]
    printed = phc5_sector_population(workbook)
    dupes = int(printed.duplicated(["district_name", "sector_name"]).sum())
    checks.append(
        _check(key, "pass" if len(printed) == 416 and not dupes else "fail",
               f"{PHC5} Tables 96-100: {len(printed)} sectors in {printed.district_name.nunique()}"
               f" districts, {dupes} repeated; Rwanda {printed.population.sum():,}")
    )  # fmt: skip
    try:
        printed["geo_id"] = sector_ids(key, printed.district_name, printed.sector_name).to_numpy()
    except SchemaError as error:
        return [*checks, _check(key, "fail", str(error))]
    staged = _staged(raw, key, "census_2022_sectors.csv")
    if staged is None or "population" not in staged:
        return [*checks, _check(key, "warn", "no staged census sector table to compare with")]
    staged["geo_id"] = sector_ids(key, staged.district_name, staged.sector_name).to_numpy()
    joined = printed.merge(staged[["geo_id", "population"]], on="geo_id", how="outer",
                           suffixes=("", "_staged"), indicator=True)  # fmt: skip
    unmatched = int((joined._merge != "both").sum())
    differ = joined[(joined._merge == "both") & (joined.population != joined.population_staged)]
    checks.append(
        _check(key, "pass" if not unmatched and differ.empty else "fail",
               f"sector populations: {len(joined) - unmatched} joined by NISR code, {unmatched} "
               f"unmatched, {len(differ)} differ from the staged census table")
    )  # fmt: skip
    staged["district_id"] = district_ids(key, staged.district_name).to_numpy()
    sums = staged.groupby("district_id")[
        [c for c in ("population", "households", "hh_improved_water") if c in staged]
    ].sum()
    tables = {t: _district_table(workbook[t]) for t in ("Table 1", "Table 58", "Table 68")}
    by_id = {}
    for table, series in tables.items():
        series = series[~series.index.str.casefold().isin([*PROVINCE_CODES, "rwanda"])]
        series = series[~series.index.str.casefold().str.contains("province")]
        try:
            by_id[table] = pd.Series(
                series.to_numpy(), index=district_ids(key, series.index.to_series()).to_numpy()
            )
        except SchemaError as error:
            checks.append(_check(key, "fail", f"{table}: {error}"))
    if "Table 1" in by_id:
        gap = (sums.population - by_id["Table 1"]).abs().max()
        checks.append(_check(key, "pass" if gap == 0 else "warn",
                             f"Table 1 district populations vs staged sector sums: largest "
                             f"difference {gap:,.0f}"))  # fmt: skip
    if "Table 58" in by_id and "households" in sums:
        gap = (sums.households - by_id["Table 58"]).abs().max()
        checks.append(_check(key, "pass" if gap == 0 else "warn",
                             f"Table 58 private households vs staged sector sums: largest "
                             f"difference {gap:,.0f}"))  # fmt: skip
    if "Table 68" in by_id and "hh_improved_water" in sums:
        share = 100 * sums.hh_improved_water / sums.households
        gap = (share - by_id["Table 68"]).abs().max()
        checks.append(_check(key, "pass" if gap <= 0.2 else "warn",
                             f"Table 68 improved drinking water (district %) vs staged "
                             f"sector counts: largest difference {gap:.1f} points "
                             "(rounding)"))  # fmt: skip
    checks.append(
        _check(key, "warn", "district level only, not staged: Table 71 (unshared improved toilet, "
               "72.1% nationally) and Table 74 (any electricity, 61.0%; the app's electricity is "
               "grid connection by sector, 52.5%). A district value is never copied to its "
               "sectors, so sanitation stays empty in the home-environment factor")
    )  # fmt: skip
    return checks


# --- NISR projection workbooks ----------------------------------------------------------------
def projection_title(path: Path) -> str | None:
    """The workbook's title if it is an NISR 2023-2032 projection table, else None."""
    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    for row in sheet.iter_rows(values_only=True, max_row=3):
        for cell in row:
            if isinstance(cell, str) and "2023-2032" in cell and "population" in cell.lower():
                return _text(cell)
    return None


def _year_columns(rows: list[tuple]) -> tuple[int, dict[int, int]]:
    """(index of the header row with years, {year: column of 'Both sexes'})."""
    for i, row in enumerate(rows[:6]):
        years = {
            cell: j for j, cell in enumerate(row) if isinstance(cell, int) and 2023 <= cell <= 2032
        }
        if len(years) == 10:
            return i, years
    raise ValueError("no row with the years 2023 to 2032")


def projection_workbooks(paths: list[Path], raw: Path = RAW) -> list:
    """Muhanga single ages and province five-year groups against the staged district table."""
    from openpyxl import load_workbook

    key = projections.DATASET.key
    staged = _staged(raw, key, projections.FILE)
    if staged is None:
        return [_check(key, "warn", "no staged projections to compare with")]
    staged["district_id"] = district_ids(key, staged.district_name).to_numpy()
    codes = pd.read_csv(REFERENCE / "district_name_crosswalk.csv", dtype=str)
    province_of = codes[codes.source == "*"].set_index("geo_id").nisr_district_code.str[:3]
    staged["province"] = staged.district_id.map(province_of)
    checks = []
    for path in paths:
        title = projection_title(path) or ""
        rows = list(
            load_workbook(path, read_only=True, data_only=True)
            .worksheets[0]
            .iter_rows(values_only=True)
        )
        try:
            header, years = _year_columns(rows)
        except ValueError as error:
            checks.append(_check(key, "fail", f"{path.name}: {error}"))
            continue
        first = years[2023]
        body = {_text(r[0]): r for r in rows[header + 2 :]
                if r[0] is not None and isinstance(r[first], int | float)}  # fmt: skip
        total = body.get("Total")
        if total is None:
            checks.append(_check(key, "fail", f"{path.name}: no Total row"))
            continue
        parts = [r for k, r in body.items() if k != "Total"]
        adds = all(sum(r[c] for r in parts) == total[c] for c in years.values())
        sexes = all(r[c] == r[c + 1] + r[c + 2] for r in body.values() for c in years.values())
        district = re.search(r"in (\w+) District", title)
        if district:
            name = district.group(1)
            mine = staged[staged.district_name == name].set_index("year")
            ages = [body[a] for a in ("3", "4", "5") if a in body]
            differ = [y for y, c in years.items() if mine.population.get(y) != total[c]
                      or mine.children_3_5.get(y) != sum(r[c] for r in ages)]  # fmt: skip
            checks.append(
                _check(key, "pass" if adds and sexes and not differ and len(ages) == 3 else "fail",
                       f"{path.name}: {name} single ages add up to the total and female + male = "
                       f"both sexes; population and children aged 3 to 5 match the staged table "
                       f"in {10 - len(differ)} of 10 years (2026: {total[years[2026]]:,} people, "
                       f"{sum(r[years[2026]] for r in ages):,} aged 3 to 5)")
            )  # fmt: skip
            continue
        province = next((c for n, c in PROVINCE_CODES.items() if n in title.lower()
                         or n.split()[0] in title.lower()), None)  # fmt: skip
        if province is None:
            checks.append(_check(key, "warn", f"{path.name}: area not recognised in '{title}'"))
            continue
        sums = staged[staged.province == province].groupby("year").population.sum()
        differ = [y for y, c in years.items() if sums.get(y) != total[c]]
        checks.append(
            _check(key, "pass" if adds and sexes and not differ else "fail",
                   f"{path.name}: {province} five-year groups add up to the total; the total "
                   f"equals the sum of its staged districts in {10 - len(differ)} of 10 years. "
                   "Five-year groups (0-4, 5-9) cannot give ages 3 to 5: verification only")
        )  # fmt: skip
    return checks


# --- MINEDUC national workbook ----------------------------------------------------------------
def is_mineduc_workbook(sheets: list[str]) -> bool:
    return set(sheets) >= MINEDUC_SHEETS


def mineduc_workbook(path: Path) -> tuple[pd.DataFrame, list]:
    """The national series in the layout intake.cross_checks reads, and its checks."""
    key = yearbook.DATASET.key
    sheets = pd.read_excel(path, sheet_name=None)
    summary, long = sheets["Annual_Summary"], sheets["Long_Data"]
    checks = []
    levels = sorted(set(long.geographic_level.astype(str)))
    checks.append(
        _check(key, "warn" if levels == ["National"] else "fail",
               f"{path.name}: {len(long)} rows, geographic level {levels}; a transcription of "
               "the 2024/25 yearbook (not a MINEDUC-issued file). National only: it checks the "
               "staged district table and is never spread to districts or sectors")
    )  # fmt: skip
    years = tuple(summary.academic_year.astype(str))
    ordered = years == tuple(sorted(years)) and len(years) == 5
    checks.append(_check(key, "pass" if ordered else "fail", f"school years {', '.join(years)}"))
    sexes = (
        summary.learners_male + summary.learners_female != summary.preprimary_learners_total
    ).sum()
    staff = (
        summary.qualified_teachers + summary.unqualified_teachers != summary.total_teaching_staff
    ).sum()
    checks.append(_check(key, "pass" if not sexes and not staff else "warn",
                         f"male + female = learners in all but {sexes} years; qualified + "
                         f"unqualified = teaching staff in all but {staff} years"))  # fmt: skip
    national = pd.concat(
        [
            pd.DataFrame(
                {
                    "indicator": name,
                    "period": summary.academic_year.astype(str),
                    "value": summary[column].astype(float),
                }
            )
            for name, column in (
                ("total_preprimary_pupils", "preprimary_learners_total"),
                ("total_preprimary_staff", "total_teaching_staff"),
                ("qualified_staff_percent", "qualified_teachers_pct"),
                ("schools_with_preprimary", "schools_total"),
                ("preprimary_classrooms", "classrooms_total"),
            )
        ],
        ignore_index=True,
    ).assign(file=path.name, period_type="school year", unit="count")
    national.loc[national.indicator == "qualified_staff_percent", "unit"] = "percent"
    return national, checks


def mineduc_capacity_checks(wide: pd.DataFrame, national: pd.DataFrame) -> list:
    """Schools and classrooms by district against the national workbook."""
    key = yearbook.DATASET.key
    checks = []
    series = national.set_index(["indicator", "period"]).value
    for measure, indicator in (("schools", "schools_with_preprimary"),
                               ("classrooms", "preprimary_classrooms")):  # fmt: skip
        for year, total in wide.groupby("school_year")[measure].sum(min_count=1).dropna().items():
            if (indicator, year) in series.index:
                gap = int(total) - int(series[(indicator, year)])
                checks.append(_check(key, "pass" if gap == 0 else "warn",
                                     f"{measure} {year}: districts {int(total):,} vs national "
                                     f"{int(series[(indicator, year)]):,}"
                                     + (f" (difference {gap:+,})" if gap else "")))  # fmt: skip
    return checks


# --- school register CSV ----------------------------------------------------------------------
def is_school_csv(columns: set[str]) -> bool:
    return columns >= SCHOOL_CSV_COLUMNS


def school_csv(path: Path, geojson: Path) -> list:
    """The CSV export against the GeoJSON the distance layer reads: same records, same points."""
    from irerero.data.sources import schools

    key = schools.DATASET.key
    table = pd.read_csv(path, dtype=str, encoding="utf-8-sig")
    points, _ = schools.read(geojson)
    table["school_code"] = table.school_code.str.replace(r"\.0$", "", regex=True)
    points["school_code"] = points.school_code.astype(str).str.replace(r"\.0$", "", regex=True)
    joined = table.merge(points[["school_code", "lon", "lat"]], on="school_code", how="outer",
                         indicator=True)  # fmt: skip
    only = joined._merge.value_counts()
    both = joined[joined._merge == "both"]
    shift = np.maximum(
        (both.geometry_longitude.astype(float) - both.lon).abs(),
        (both.geometry_latitude.astype(float) - both.lat).abs(),
    ).max()
    return [
        _check(key, "pass" if only.get("left_only", 0) == only.get("right_only", 0) == 0
               else "fail",
               f"{path.name}: {len(table)} records, {table.school_code.duplicated().sum()} "
               f"repeated codes; {only.get('both', 0)} share a code with the GeoJSON, "
               f"{only.get('left_only', 0)} only in the CSV, {only.get('right_only', 0)} only in "
               "the GeoJSON"),
        _check(key, "pass" if shift < 1e-6 else "warn",
               f"CSV coordinates vs GeoJSON points: largest difference {shift:.1e} degrees. The "
               "two files are the same records; only the GeoJSON is staged, so no school is "
               "counted twice"),
    ]  # fmt: skip


# --- WorldPop rasters -------------------------------------------------------------------------
def worldpop_layers(paths: list[Path]) -> list:
    """Which layers arrived, their grid and totals; staged only when a full under-5 set is in."""
    import rasterio

    key = worldpop.DATASET.key
    checks = []
    sets: dict[tuple[str, str], set[tuple[str, str]]] = {}
    for path in sorted(paths):
        match = WORLDPOP_NAME.match(path.name)
        if not match:
            checks.append(_check(key, "warn", f"{path.name}: not a recognised WorldPop name"))
            continue
        with rasterio.open(path) as source:
            band = source.read(1, masked=True)
            checks.append(
                _check(key, "pass" if source.crs and source.crs.to_epsg() == 4326 else "warn",
                       f"{path.name}: {source.crs}, {source.res[0] * 111_320:.0f} m cells, "
                       f"{source.width}x{source.height}, {band.count():,} populated cells, total "
                       f"{float(band.sum()):,.0f}, {int((band < 0).sum())} negative")
            )  # fmt: skip
        sets.setdefault((match["year"], match["product"]), set()).add((match["sex"], match["age"]))
    for (year, product), layers in sorted(sets.items()):
        ages = ("0", "1") if product == "constrained_UNadj" else ("00", "01")
        need = {(s, a) for s in "fm" for a in ages}
        missing = sorted(need - layers)
        names = [f"rwa_{s}_{a}_{year}_{product}.tif" for s, a in missing]
        checks.append(
            _check(key, "warn" if missing else "pass",
                   f"{year} ({product}): under-5 needs female and male, ages under 1 and 1 to 4; "
                   + (f"missing {', '.join(names)}, so this set is not staged and the app keeps "
                      "its census-share stand-in (one sex alone would halve and skew the count)"
                      if missing else "complete"))
        )  # fmt: skip
        other = sorted(f"{s}_{a}" for s, a in layers if a not in UNDER5)
        if other:
            checks.append(_check(key, "warn", f"{year}: {', '.join(other)} are not under-5 "
                                 "layers (ages 5 to 9, or all under 18): not used"))  # fmt: skip
    return checks
